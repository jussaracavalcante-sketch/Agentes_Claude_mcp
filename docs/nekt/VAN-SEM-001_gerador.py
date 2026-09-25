# -*- coding: utf-8 -*-
"""VAN-SEM-001 - Levantamento para a Camada Semantica, por setor.
Tres blocos por aba: perguntas de decisao, indicadores e vocabulario.
Linhas pre-preenchidas vem de evidencia ja medida na plataforma - o setor valida."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

COD, VER, DATA = "VAN-SEM-001", "1.0", "11/09/2026"
AZUL, AZUL2, CINZA, AMAREL, VERDE, LILAS = "1F2D5C","2E75B6","F2F2F2","FFF3CD","E8F3EC","EEF2FA"
B = Border(*[Side(style="thin", color="BFBFBF")]*4)
F_TIT=Font(name="Arial",size=14,bold=True,color=AZUL)
F_SUB=Font(name="Arial",size=9,color="595959")
F_HDR=Font(name="Arial",size=9,bold=True,color="FFFFFF")
F_BAND=Font(name="Arial",size=11,bold=True,color="FFFFFF")
F_TXT=Font(name="Arial",size=10)
F_PRE=Font(name="Arial",size=9,color="1F3864")
F_LBL=Font(name="Arial",size=10,bold=True)
F_CALC=Font(name="Arial",size=10,bold=True,color="1F4E2C")
PRE="pré-preenchido — confirme ou corrija"

# ---- LINHAS: (aba, nome, [perguntas], [indicadores], [termos]) -------------
# pergunta: (pergunta, decisao que depende, frequencia, onde busca hoje)
# indicador: (nome, como se calcula, unidade, recorte, o que NAO vale fazer)
# termo: (termo, significado, onde aparece no dado)
S = []
S.append(("Mídia Paga","Mídia Paga",
 [("Quanto cada cliente investiu no período, por plataforma?","Realocação de verba entre contas","Semanal","rfn_midia__desempenho_diario"),
  ("Quais campanhas estão sem entrega e há quantos dias?","Pausar ou reativar campanha","Semanal","rfn_midia__desempenho_diario"),
  ("Qual o custo por conversão por conta e por canal?","Ajuste de lance e de canal","Semanal","rfn_midia__desempenho_diario")],
 [("Investimento","Somar investimento_micros e dividir por 1.000.000","R$ ou US$","Conta, campanha, dia",
   "NÃO somar a coluna investimento (é arredondada por linha: acumulou R$ 0,90 em 42 mil linhas). NÃO somar BRL com USD."),
  ("Cobertura demográfica","Verba não-PMax ÷ verba total","%","Conta",
   "NÃO usar as tabelas de faixa etária/gênero para totalizar verba: PERFORMANCE_MAX não publica breakdown demográfico."),
  ("Cobertura de termo de busca","Verba com termo publicado ÷ verba de SEARCH+SHOPPING","%","Conta",
   "NÃO usar a verba total como denominador: DISPLAY, VIDEO e PMax não têm termo de busca para publicar."),
  ("Conta defasada","Dias desde a última data com entrega","dias","Conta",
   "NÃO ler como fonte parada: em 11 de 36 contas a extração rodou e teve sucesso — a conta é que parou de anunciar.")],
 [("Verba","Investimento veiculado na plataforma de mídia, em micros na origem","rfn_midia__desempenho_diario.investimento_micros"),
  ("Conta","Unidade de cobrança da plataforma. Contas de nome parecido NÃO se fundem (R-003).","id_conta"),
  ("PMax","Campanha PERFORMANCE_MAX: não publica desempenho por anúncio nem breakdown demográfico","canal / subcanal"),
  ("Grão misto","Linha de anúncio somada a linha de campanha só onde não há anúncio","grao_origem")]))
S.append(("Inbound","Inbound",
 [("Quantas conversões vieram de cada origem de tráfego?","Onde concentrar esforço de geração","Semanal","rfn_marketing__conversao"),
  ("Que parcela das conversões liga a uma campanha de Google Ads?","Atribuição de investimento a resultado","Mensal","rfn_marketing__conversao")],
 [("Conversões por origem","Contar id_conversao agrupando por origem_canonica","un.","Cliente, canal, mês",
   "NÃO ler só utm_source: ele traz 1.400 linhas com um único valor distinto e dá a falsa impressão de que não há origem."),
  ("Cobertura de funil","Conversões com gad_campaignid ÷ total de conversões","%","Cliente",
   "Cobertura baixa (5,2%, 908 de 17.602) é o número certo, não erro. NÃO casar por nome de cliente para 'melhorar': só 8 rótulos batem e a R-003 proíbe fundir os casos 3-para-1.")],
 [("Origem canônica","Origem de tráfego já decodificada dos 5 formatos brutos","origem_canonica"),
  ("Conversão","Um evento de contato do RD Station","id_conversao"),
  ("Canal pago","Marcação booleana de origem paga","canal_pago")]))
S.append(("Account","Account (Atendimento)",
 [("Qual o volume de jobs por cliente e em que etapa estão?","Priorização e alerta de prazo","Semanal","iClips — fato_atividade"),
  ("Qual o faturamento por cliente e o que está por faturar?","Conversa comercial e cobrança","Mensal","Conta Azul — vw_funil_faturamento")],
 [("Jobs por cliente","Contar job_id distintos por client_name","un.","Cliente, mês","NÃO usar client_name para agrupar grupo econômico — use client_group_name ou CNPJ."),
  ("Valor faturado","Somar valor_faturado_cliente","R$","Cliente, competência",
   "Para BRUTO × LÍQUIDO use tipo_faturamento (fonte autoritativa do ERP), não a inferência por CNPJ da nota.")],
 [("PI","Pedido de Inserção — o contrato de veiculação de mídia","silver_pi_insercao.pi"),
  ("Job","Unidade de trabalho da agência","fato_atividade.job_id"),
  ("Grupo econômico","Conjunto de clientes do mesmo dono, com CNPJs distintos","client_group_name / tags_id")]))
S.append(("Criação","Criação",
 [("Quantas peças foram entregues por etapa de workflow?","Dimensionamento de equipe","Semanal","iClips — fato_atividade"),
  ("Qual o índice de refação e onde ele se concentra?","Ajuste de briefing e de processo","Mensal","iClips — fato_atividade")],
 [("Peças por etapa","Contar piece_id por workflow_normalizado","un.","Cliente, mês, responsável",
   "NÃO agrupar pela coluna workflow crua: grafias mistas fragmentaram 29.765 atividades (VAN-73). Use workflow_normalizado."),
  ("Índice de refação","Peças com refacao ÷ total de peças","%","Cliente, responsável","Confirmar com a equipe o que conta como refação antes de publicar o indicador.")],
 [("Peça","Entregável individual dentro de um job","fato_atividade.piece_id"),
  ("Workflow","Etapa do fluxo da peça — é o nome da etapa, não do template","workflow_normalizado"),
  ("Refação","Retrabalho de peça já entregue","fato_atividade.refacao")]))
S.append(("Direção de Arte","Direção de Arte",
 [("Qual o tempo médio entre entrada e aprovação da peça?","Negociação de prazo com Account","Mensal","iClips — fato_atividade")],
 [("Tempo de ciclo da peça","data_fim_peca_workflow menos data_inicio_peca_workflow","dias","Cliente, tipo de peça",
   "Conferir se estimated_time e sla_time estão preenchidos: linha sem eles não entra na média e enviesa o resultado.")],
 [("Aprovação","Marco em que a peça é aceita pelo cliente","project_approval_date")]))
S.append(("Dir. Operações","Diretoria de Operações",
 [("Quantos jobs estão abertos, por responsável e por prazo?","Redistribuição de carga","Semanal","rfn_operacao__job (VJOB)")],
 [("Jobs em aberto","Contar jobs com status diferente de concluído","un.","Responsável, cliente",
   "O CONTADOR DE ATRASO DO VJOB ESTÁ FORMALMENTE CLASSIFICADO COMO INVÁLIDO pela própria casa enquanto o saneamento de prazos não conclui. NÃO usar como KPI.")],
 [("Job","Tarefa corporativa registrada no VJOB","rfn_operacao__job"),
  ("Prazo inconsistente","Tarefa com prazo anterior à própria publicação","marca do PMO")]))
S.append(("RH","Recursos Humanos",
 [("Qual o turnover do período e onde ele se concentra?","Ação de retenção por área","Mensal","silver_colaborador_rel"),
  ("Qual o tempo médio de casa por área?","Plano de carreira e sucessão","Trimestral","silver_colaborador_rel")],
 [("Turnover","Desligamentos no período ÷ headcount médio","%","Área, departamento, nível",
   "NÃO usar data_desligamento crua em série: use data_desligamento_efetiva. A crua guarda resíduo de vínculo antigo em recontratação."),
  ("Headcount","Contar pessoas com situacao = ATIVO","un.","Área, departamento",
   "Tratar ATIVO_INCONSISTENTE à parte: é cadastro a corrigir na origem, não uma categoria de pessoa.")],
 [("Situação","Estado do vínculo: ATIVO, AVISO_PREVIO, DESLIGADO ou ATIVO_INCONSISTENTE","colaborador_rel.situacao"),
  ("Tempo de casa","Meses entre admissão e hoje, ou até o desligamento efetivo","tempo_casa_meses")]))
S.append(("Dir. Executiva","Diretoria Executiva",
 [("Qual a receita por cliente e por linha de serviço?","Decisão de portfólio","Mensal","gold_vw_fin_cliente"),
  ("Qual o caixa realizado contra o projetado?","Decisão de investimento","Mensal","gold_vw_fato_dfc_dia")],
 [("Receita por cliente","Somar valor agrupando por cliente_key","R$","Cliente, grupo, linha de serviço",
   "SEMPRE separar por flag_projecao: misturar realizado com projeção produz número que não existe."),
  ("MRR","Somar mrr do mês","R$","Cliente, produto, família","Distinguir tipo = plano de tipo = serviço: são chaves de produto diferentes (ORI-603).")],
 [("Realizado × Projetado","Separação obrigatória em toda leitura financeira","flag_projecao"),
  ("Grupo","Agrupamento econômico acima do cliente","grupo_key / grupo_nome")]))
S.append(("Social Media","Social Media",
 [("Quantas peças de social foram publicadas por cliente?","Cumprimento de contrato","Mensal","iClips — fato_atividade")],
 [("Peças publicadas","Contar peças com prefixo [SOCIAL MEDIA] no catálogo","un.","Cliente, mês",
   "Separar peça produzida de peça veiculada: são coisas diferentes e o dado hoje registra a produção.")],
 [("Peça de social","Entregável do catálogo com prefixo [SOCIAL MEDIA]","dim_peca_canonica.prefixo")]))
S.append(("Financeiro","Financeiro / Controladoria",
 [("Qual a inadimplência por cliente e por faixa de atraso?","Régua de cobrança","Semanal","vw_inad_titulos_vbot"),
  ("Qual o fluxo de caixa diário por conta?","Programação de pagamento","Diária","vw_fluxo_caixa_diario")],
 [("Inadimplência","Somar valor_aberto de títulos vencidos","R$","Cliente, faixa de atraso",
   "Títulos com flag_migrado vieram do iClips em 01/06/2026 e a baixa ocorre no Conta Azul — não contar cobrança em dobro."),
  ("Perda (write-off)","Somar o campo perda","R$","Cliente, período",
   "A API devolve perda como objeto {data, valor}, não escalar — ler perda.valor, senão o total sai zerado.")],
 [("Parcela","Unidade de cobrança do evento financeiro","ca_fato_evento_financeiro.id_parcela"),
  ("Baixa","Registro de pagamento da parcela","campo baixas"),
  ("Conciliado","Parcela batida com o extrato","conciliado")]))
S.append(("Planej. Estrat. e Inovação","Planejamento, Estratégia e Inovação",
 [("Onde o mesmo termo significa coisas diferentes entre setores?","Padronização do vocabulário da casa","Trimestral","este próprio levantamento"),
  ("Que perguntas de negócio nenhum setor consegue responder hoje?","Priorização do roadmap de dados","Trimestral","coluna 'Hoje consegue responder?' das outras abas")],
 [("Perguntas sem resposta","Contar perguntas marcadas 'Não' em todas as abas","un.","Setor",
   "Não é indicador de desempenho de setor — é fila de trabalho da plataforma de dados.")],
 [("Camada semântica","Documentação que diz o que o número significa e como lê-lo sem errar","Nekt — documentos de contexto"),
  ("Refined","Camada oficial de consumo, onde moram as regras de negócio","prefixo rfn_")]))

LIST={"Freq":["Diária","Semanal","Quinzenal","Mensal","Trimestral","Semestral","Anual","Sob demanda"],
      "Consegue":["Sim","Parcial","Não"],
      "Setores":[x[1] for x in S]}
COLW=[40,46,16,22,40,22,26]
R_FICHA, R_B1, R_B2, R_B3 = 7, 14, 30, 51
N1, N2, N3 = 13, 18, 13     # linhas de cada bloco
wb=Workbook(); wb.remove(wb.active)

def band(ws,row,txt):
    ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=7)
    c=ws.cell(row=row,column=1,value=txt); c.font=F_BAND
    c.fill=PatternFill("solid",fgColor=AZUL2); c.alignment=Alignment(vertical="center",indent=1)
    ws.row_dimensions[row].height=22

def hdr(ws,row,nomes):
    for i,n in enumerate(nomes,start=1):
        c=ws.cell(row=row,column=i,value=n)
        if n:
            c.font=F_HDR; c.fill=PatternFill("solid",fgColor=AZUL)
        c.border=B; c.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center")
    ws.row_dimensions[row].height=34

def bloco(ws,row0,n,cols_usadas):
    for r in range(row0,row0+n):
        for c in range(1,8):
            cel=ws.cell(row=r,column=c); cel.border=B; cel.font=F_TXT
            cel.alignment=Alignment(wrap_text=True,vertical="top")
            if c not in cols_usadas: cel.fill=PatternFill("solid",fgColor="FAFAFA")
            elif (r-row0)%2==1: cel.fill=PatternFill("solid",fgColor=CINZA)
        ws.row_dimensions[r].height=34

for aba,nome,PERG,IND,TERM in S:
    ws=wb.create_sheet(aba)
    for i,w in enumerate(COLW,start=1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.merge_cells("A1:G1"); ws.cell(row=1,column=1,value=f"Camada Semântica — {nome}").font=F_TIT
    ws.merge_cells("A2:G2"); ws.cell(row=2,column=1,
        value=f"{COD}  |  Versão {VER}  |  Emissão {DATA}  |  Plataforma de Dados Nekt  |  uso interno").font=F_SUB
    ws.merge_cells("A3:G3"); ws.cell(row=3,column=1,
        value="Objetivo: registrar o que o número significa neste setor e como lê-lo sem errar. Não é mapeamento de processo.").font=F_SUB
    ws.merge_cells("A5:G5")
    lg=ws.cell(row=5,column=1,value="COMO PREENCHER: as linhas em azul-claro já vêm preenchidas com o que a plataforma de dados "
        "hoje sabe deste setor — sua tarefa é CONFIRMAR ou CORRIGIR, e acrescentar o que falta nas linhas em branco. "
        "Se discordar de uma linha pré-preenchida, sobrescreva: a sua versão vale mais que a minha.")
    lg.font=Font(name="Arial",size=9,bold=True,color="1F4E2C"); lg.fill=PatternFill("solid",fgColor=VERDE)
    lg.alignment=Alignment(wrap_text=True,vertical="center"); ws.row_dimensions[5].height=34

    band(ws,R_FICHA,"FICHA DO SETOR")
    fichas=[("Quem responde por este preenchimento",""),("Cargo / função",""),
            ("Sistemas que o setor usa no dia a dia",""),("Com que frequência o setor decide com dado",""),
            ("Período que costuma olhar (ex.: mês corrente, últimos 90 dias)","")]
    for k,(lab,val) in enumerate(fichas):
        r=R_FICHA+1+k
        c=ws.cell(row=r,column=1,value=lab); c.font=F_LBL; c.border=B
        c.alignment=Alignment(wrap_text=True,vertical="center")
        ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=7)
        v=ws.cell(row=r,column=2,value=val); v.fill=PatternFill("solid",fgColor=AMAREL); v.border=B
        ws.row_dimensions[r].height=20

    band(ws,R_B1,"1. PERGUNTAS QUE O SETOR PRECISA RESPONDER COM DADO   —   define o que a camada Refined tem de servir")
    hdr(ws,R_B1+1,["Pergunta","Decisão que depende dela","Frequência","Hoje consegue responder?","Onde busca hoje","","Origem da linha"])
    bloco(ws,R_B1+2,N1,{1,2,3,4,5,7})
    for k,(p,d,f,o) in enumerate(PERG):
        r=R_B1+2+k
        for col,val in [(1,p),(2,d),(3,f),(5,o),(7,PRE)]:
            c=ws.cell(row=r,column=col,value=val); c.font=F_PRE
            c.fill=PatternFill("solid",fgColor=LILAS)
        for col in (4,6):
            ws.cell(row=r,column=col).fill=PatternFill("solid",fgColor=AMAREL)

    band(ws,R_B2,"2. INDICADORES   —   nome, cálculo e, sobretudo, o que NÃO vale fazer com ele")
    hdr(ws,R_B2+1,["Indicador","Como se calcula (em palavras)","Unidade","Recorte habitual","O que NÃO vale fazer com ele","Dono da definição","Origem da linha"])
    bloco(ws,R_B2+2,N2,{1,2,3,4,5,6,7})
    for k,(n,calc,u,rec,nao) in enumerate(IND):
        r=R_B2+2+k
        for col,val in [(1,n),(2,calc),(3,u),(4,rec),(5,nao),(7,PRE)]:
            c=ws.cell(row=r,column=col,value=val); c.font=F_PRE
            c.fill=PatternFill("solid",fgColor=LILAS)
        ws.cell(row=r,column=6).fill=PatternFill("solid",fgColor=AMAREL)
        ws.row_dimensions[r].height=48

    band(ws,R_B3,"3. VOCABULÁRIO DO SETOR   —   o bloco mais barato e o de maior retorno")
    hdr(ws,R_B3+1,["Termo","O que significa NESTE setor","","","Onde aparece no dado","","Origem da linha"])
    bloco(ws,R_B3+2,N3,{1,2,5,7})
    for k,(t,sig,onde) in enumerate(TERM):
        r=R_B3+2+k
        for col,val in [(1,t),(2,sig),(5,onde),(7,PRE)]:
            c=ws.cell(row=r,column=col,value=val); c.font=F_PRE
            c.fill=PatternFill("solid",fgColor=LILAS)

    for lista,col,r0,n in [("Freq","C",R_B1+2,N1),("Consegue","D",R_B1+2,N1),("Freq","C",R_B2+2,N2)]:
        pass
    d1=DataValidation(type="list",formula1='"'+",".join(LIST["Freq"])+'"',allow_blank=True)
    d2=DataValidation(type="list",formula1='"'+",".join(LIST["Consegue"])+'"',allow_blank=True)
    ws.add_data_validation(d1); ws.add_data_validation(d2)
    d1.add(f"C{R_B1+2}:C{R_B1+1+N1}"); d2.add(f"D{R_B1+2}:D{R_B1+1+N1}")
    ws.freeze_panes="A7"; ws.sheet_view.showGridLines=False

print("setores:",len(S))

# ---------------- CONSOLIDADO ----------------
CC=[("Setor",34),("Ficha (de 5)",12),("Perguntas",12),("Sem resposta hoje",16),
    ("Indicadores",12),("Com regra de 'não vale'",18),("Termos",10),("Situação",22)]
ws=wb.create_sheet("Consolidado")
for i,(n,w) in enumerate(CC,start=1): ws.column_dimensions[get_column_letter(i)].width=w
ws.merge_cells("A1:H1"); ws.cell(row=1,column=1,value="Consolidado do Levantamento Semântico").font=F_TIT
ws.merge_cells("A2:H2"); ws.cell(row=2,column=1,value=f"{COD}  |  Versão {VER}  |  Emissão {DATA}").font=F_SUB
ws.merge_cells("A4:H4")
lg=ws.cell(row=4,column=1,value="ESTA ABA É TODA CALCULADA — não digite nada aqui. "
  "'Situação' fica Pronto quando o setor tem a ficha completa, ao menos 2 perguntas, 3 indicadores e 3 termos. "
  "A coluna 'Sem resposta hoje' é a fila de trabalho da plataforma: são perguntas de negócio que o dado ainda não responde.")
lg.font=Font(name="Arial",size=9,bold=True,color="1F4E2C"); lg.fill=PatternFill("solid",fgColor=VERDE)
lg.alignment=Alignment(wrap_text=True,vertical="center"); ws.row_dimensions[4].height=34
hdr(ws,6,[n for n,_ in CC])
for k,(aba,nome,_,_,_) in enumerate(S):
    r=7+k; q=f"'{aba}'"
    c=ws.cell(row=r,column=1,value=nome); c.font=F_LBL
    ws.cell(row=r,column=2,value=f'=COUNTIF({q}!$B$8:$B$12,"?*")')
    ws.cell(row=r,column=3,value=f'=COUNTIF({q}!$A$16:$A$28,"?*")')
    ws.cell(row=r,column=4,value=f'=COUNTIF({q}!$D$16:$D$28,"Não")')
    ws.cell(row=r,column=5,value=f'=COUNTIF({q}!$A$32:$A$49,"?*")')
    ws.cell(row=r,column=6,value=f'=COUNTIF({q}!$E$32:$E$49,"?*")')
    ws.cell(row=r,column=7,value=f'=COUNTIF({q}!$A$53:$A$65,"?*")')
    ws.cell(row=r,column=8,value=f'=IF(AND($B{r}=5,$C{r}>=2,$E{r}>=3,$G{r}>=3),"Pronto","Falta preencher")')
    for c2 in range(1,9):
        cel=ws.cell(row=r,column=c2); cel.border=B
        cel.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center" if c2>1 else "left")
        if c2>1: cel.font=F_CALC; cel.fill=PatternFill("solid",fgColor=VERDE)
    ws.row_dimensions[r].height=26
rt=7+len(S)
ws.cell(row=rt,column=1,value="TOTAL").font=F_LBL
for c2 in range(2,8):
    L=get_column_letter(c2); ws.cell(row=rt,column=c2,value=f"=SUM({L}7:{L}{rt-1})").font=F_LBL
ws.cell(row=rt,column=8,value=f'=COUNTIF($H$7:$H{rt-1},"Pronto")&" de {len(S)} setores"').font=F_LBL
for c2 in range(1,9):
    cel=ws.cell(row=rt,column=c2); cel.border=B; cel.fill=PatternFill("solid",fgColor=CINZA)
    cel.alignment=Alignment(horizontal="center" if c2>1 else "left")
ws.sheet_view.showGridLines=False

# ---------------- INSTRUÇÕES ----------------
ws=wb.create_sheet("Instruções",0)
ws.column_dimensions["A"].width=4; ws.column_dimensions["B"].width=32; ws.column_dimensions["C"].width=98
ws.sheet_view.showGridLines=False
ws.cell(row=1,column=2,value="Levantamento para a Camada Semântica").font=Font(name="Arial",size=16,bold=True,color=AZUL)
ws.cell(row=2,column=2,value=f"{COD}  |  Versão {VER}  |  Emissão {DATA}  |  uso interno").font=F_SUB
ws.cell(row=3,column=2,value="Plataforma de Dados Nekt — Vanguarda Martech").font=F_SUB
BL=[("O que é camada semântica","É a documentação que diz O QUE O NÚMERO SIGNIFICA e COMO LÊ-LO SEM ERRAR. "
  "Não descreve como o trabalho é feito — descreve como o dado deve ser interpretado. "
  "Sem ela, cada relatório recalcula a métrica do seu jeito e dois times chegam a números diferentes para a mesma pergunta."),
 ("O que este formulário NÃO é","Não é mapeamento de processo. Não pede RACI, risco, competência, treinamento, "
  "retenção documental nem fluxo. Isso pertence ao SGQ e tem formulário próprio (VAN-FOR-SGQ-001)."),
 ("Quanto tempo leva","Entre 20 e 40 minutos por setor. A maior parte já vem preenchida; você valida."),
 ("As linhas azuis já vêm prontas","São o que a plataforma de dados hoje sabe do seu setor: indicadores que já existem, "
  "armadilhas já medidas, tabelas que já rodam. CONFIRME, CORRIJA ou APAGUE. Se você discordar, a sua versão vale — "
  "você conhece o setor, eu conheço a tabela."),
 ("Os três blocos","1) PERGUNTAS que você precisa responder com dado. É o que define o que a plataforma tem de servir; "
  "marcar 'Não' em 'Hoje consegue responder?' não é confissão de falha, é pedido de trabalho.\n"
  "2) INDICADORES: nome, cálculo em palavras, unidade, recorte.\n"
  "3) VOCABULÁRIO: o que cada termo significa no seu setor."),
 ("A coluna mais importante","'O QUE NÃO VALE FAZER COM ELE', no bloco 2. É ela que evita o erro caro. "
  "Exemplos reais desta casa: não somar investimento entre moedas diferentes; não somar a coluna já arredondada; "
  "não usar o contador de atraso do VJOB como KPI enquanto a própria casa o classifica como inválido. "
  "Se você sabe de um jeito errado de ler um número seu, escreva ali — é a informação mais valiosa do formulário."),
 ("Por que vocabulário importa tanto","'Job', 'verba', 'entrega' e 'cliente' significam coisas diferentes em Account, "
  "Mídia e Operações. Ninguém percebe até o número não bater numa reunião. Escrever a definição custa dois minutos "
  "e evita a discussão inteira."),
 ("Cores","Azul-claro = já preenchido, para validar. Amarelo = campo seu, em branco. "
  "Verde = calculado, não digite. Cinza claro = coluna não usada naquele bloco."),
 ("O que acontece depois","Cada aba devolvida vira um documento de contexto na camada semântica da Nekt, "
  "com as fórmulas e as regras de leitura que você declarou. A partir daí, toda consulta — feita por pessoa ou por IA — "
  "passa a enxergar a sua definição antes de calcular."),
 ("Devolução","Devolver a planilha inteira, com todas as abas. A aba Consolidado mostra sozinha quem já concluiu."),
]
r=5
for t,x in BL:
    c=ws.cell(row=r,column=2,value=t); c.font=F_HDR; c.fill=PatternFill("solid",fgColor=AZUL)
    c.border=B; c.alignment=Alignment(vertical="center",wrap_text=True)
    v=ws.cell(row=r,column=3,value=x); v.font=F_TXT; v.border=B
    v.alignment=Alignment(wrap_text=True,vertical="top")
    ws.row_dimensions[r].height=max(30, 14*(x.count("\n")+1)+12.5*(len(x)//97)); r+=1

ORD=["Instruções"]+[x[0] for x in S]+["Consolidado"]
wb._sheets=sorted(wb._sheets,key=lambda s:ORD.index(s.title))
out="/home/user/Agentes_Claude_mcp/docs/nekt/VAN-SEM-001_Camada_Semantica_por_Setor.xlsx"
wb.save(out)
print("abas:",len(wb.sheetnames)); print(wb.sheetnames); print("salvo:",out)
